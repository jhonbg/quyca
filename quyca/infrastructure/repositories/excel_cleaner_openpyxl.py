from __future__ import annotations

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

from quyca.domain.repositories.excel_cleaner_interface import IExcelCleaner


class ExcelCleanerOpenpyxl(IExcelCleaner):
    """Removes validation columns and formatting from Excel files."""

    def __init__(self) -> None:
        self._targets = {"estado_de_validacion", "observacion"}
        self._no_fill = PatternFill(fill_type=None)

    def _norm(self, value: object) -> str:
        """Normalizes header values for matching."""
        if value is None:
            return ""
        return (
            str(value)
            .strip()
            .lower()
            .replace(" ", "_")
            .replace("validación", "validacion")
            .replace("observación", "observacion")
        )

    def clean(self, filepath: str) -> None:
        """Deletes target columns and clears cell fills in the workbook."""
        wb = load_workbook(filepath)

        for ws in wb.worksheets:
            cols_to_delete: list[int] = []
            for col_idx in range(1, ws.max_column + 1):
                header = ws.cell(row=1, column=col_idx).value
                if self._norm(header) in self._targets:
                    cols_to_delete.append(col_idx)

            for col_idx in reversed(cols_to_delete):
                ws.delete_cols(col_idx, 1)

            for cell in ws._cells.values():
                cell.fill = self._no_fill

        wb.save(filepath)
